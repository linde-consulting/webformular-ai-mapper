"""Import von Formularfeld-Tabellen (CSV/XLSX) - Alternative zur manuellen JSON.

Formcycle bietet keinen dokumentierten Rohexport der Feld-Definitionen (Typ,
Pflichtfeld, Bezug) als einfaches CSV - nur einen Submission-Datenexport (Spalten
= Feld-Label, ohne Metadaten) und einen komplexen internen Formular-Export
(undokumentiert, siehe README). Deshalb definiert dieses Modul einen eigenen,
einfachen Tabellen-Vertrag, den ein Kunde manuell aus Formcycle ableiten oder
direkt so pflegen kann (siehe examples/formular-vorlage.csv).

Spalten (Groß-/Kleinschreibung egal, gängige Synonyme werden erkannt):
  Abschnitt | Feld-ID | Label | Typ | Pflichtfeld | Optionen | Hilfetext

Nur "Label" und "Typ" sind zwingend. Fehlt "Feld-ID", wird sie aus dem Label
abgeleitet (Slug, bei Kollision mit laufender Nummer eindeutig gemacht).
"""
from __future__ import annotations

import csv
import io
import json
import re
from typing import Dict, List, Optional

from .formcycle_import import formular_titel, ist_formcycle_export, parse_formcycle_export
from .models import FeldTyp, FormularInput, Formularfeld

SPALTEN_ALIASE: Dict[str, List[str]] = {
    "abschnitt": ["abschnitt", "gruppe", "section"],
    "id": ["feld-id", "feldid", "id"],
    "label": ["label", "bezeichnung", "feldname", "name"],
    "typ": ["typ", "type", "feldtyp"],
    "pflichtfeld": ["pflichtfeld", "pflicht", "required"],
    "optionen": ["optionen", "auswahlwerte", "werte"],
    "hilfetext": ["hilfetext", "hilfe", "beschreibung"],
}

PFLICHT_WAHR_WERTE = {"ja", "true", "1", "x", "yes", "wahr"}

ERLAUBTE_TYPEN = FeldTyp.__args__  # type: ignore[attr-defined]


class FormularImportError(ValueError):
    pass


def _slugify(text: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", text.strip().lower()).strip("-")
    return slug or "feld"


def _finde_spalte(header: List[str], schluessel: str) -> Optional[int]:
    normalisiert = [h.strip().lower() for h in header]
    for alias in SPALTEN_ALIASE[schluessel]:
        if alias in normalisiert:
            return normalisiert.index(alias)
    return None


def _lies_zeilen_csv(inhalt: bytes) -> List[List[str]]:
    text = inhalt.decode("utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel  # Fallback: Komma-getrennt (z. B. bei nur einer Spalte)
    reader = csv.reader(io.StringIO(text), dialect)
    return [row for row in reader if any(cell.strip() for cell in row)]


def _lies_zeilen_xlsx(inhalt: bytes) -> List[List[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(inhalt), read_only=True, data_only=True)
    sheet = wb.active
    zeilen = []
    for row in sheet.iter_rows(values_only=True):
        werte = ["" if v is None else str(v) for v in row]
        if any(cell.strip() for cell in werte):
            zeilen.append(werte)
    return zeilen


def parse_feld_datei(inhalt: bytes, dateiname: str) -> List[Formularfeld]:
    dateiname_lower = dateiname.lower()
    if dateiname_lower.endswith(".xlsx"):
        zeilen = _lies_zeilen_xlsx(inhalt)
    elif dateiname_lower.endswith(".csv"):
        zeilen = _lies_zeilen_csv(inhalt)
    else:
        raise FormularImportError(
            f"Nicht unterstütztes Dateiformat: '{dateiname}'. Erlaubt sind .csv und .xlsx."
        )

    if not zeilen:
        raise FormularImportError("Die Datei enthält keine Zeilen.")

    header, *datenzeilen = zeilen
    idx_abschnitt = _finde_spalte(header, "abschnitt")
    idx_id = _finde_spalte(header, "id")
    idx_label = _finde_spalte(header, "label")
    idx_typ = _finde_spalte(header, "typ")
    idx_pflicht = _finde_spalte(header, "pflichtfeld")
    idx_optionen = _finde_spalte(header, "optionen")
    idx_hilfetext = _finde_spalte(header, "hilfetext")

    if idx_label is None or idx_typ is None:
        raise FormularImportError(
            "Pflichtspalten fehlen: 'Label' und 'Typ' müssen als Spaltenüberschrift vorhanden sein. "
            f"Gefundene Spalten: {header}"
        )

    def _wert(row: List[str], idx: Optional[int]) -> str:
        if idx is None or idx >= len(row):
            return ""
        return (row[idx] or "").strip()

    felder: List[Formularfeld] = []
    vergebene_ids: Dict[str, int] = {}

    for zeilen_nr, row in enumerate(datenzeilen, start=2):
        label = _wert(row, idx_label)
        typ = _wert(row, idx_typ).lower()
        if not label:
            continue  # leere Zeile

        if typ not in ERLAUBTE_TYPEN:
            raise FormularImportError(
                f"Zeile {zeilen_nr}: Ungültiger Typ '{typ}' für Feld '{label}'. "
                f"Erlaubt sind: {', '.join(ERLAUBTE_TYPEN)}."
            )

        feld_id = _wert(row, idx_id) or _slugify(label)
        if feld_id in vergebene_ids:
            vergebene_ids[feld_id] += 1
            feld_id = f"{feld_id}-{vergebene_ids[feld_id]}"
        else:
            vergebene_ids[feld_id] = 1

        optionen_text = _wert(row, idx_optionen)
        optionen = [o.strip() for o in optionen_text.split(";") if o.strip()] or None

        felder.append(
            Formularfeld(
                id=feld_id,
                label=label,
                typ=typ,  # type: ignore[arg-type]
                pflichtfeld=_wert(row, idx_pflicht).lower() in PFLICHT_WAHR_WERTE,
                optionen=optionen,
                abschnitt=_wert(row, idx_abschnitt) or None,
                hilfetext=_wert(row, idx_hilfetext) or None,
            )
        )

    if not felder:
        raise FormularImportError("Keine gültigen Feldzeilen gefunden.")

    return felder


def build_formular_input(
    inhalt: bytes,
    dateiname: str,
    titel: Optional[str] = None,
    leika_nummer: Optional[str] = None,
) -> FormularInput:
    """Erkennt das Dateiformat und baut daraus ein FormularInput.

    - `.json` (eigenes Format): vollständiges FormularInput (titel/leika_nummer/felder enthalten).
    - `.json` (echter FORMCYCLE-Export, "Formular exportieren"): Feldtypen/Labels werden aus
      `items` extrahiert; `titel` wird aus dem Formular übernommen, falls nicht angegeben.
    - `.csv` / `.xlsx`: reine Feld-Tabelle - `titel` muss zusätzlich übergeben werden.
    """
    dateiname_lower = dateiname.lower()

    if dateiname_lower.endswith(".json"):
        rohdaten = json.loads(inhalt)
        if ist_formcycle_export(rohdaten):
            endgueltiger_titel = titel or formular_titel(rohdaten)
            if not endgueltiger_titel:
                raise FormularImportError(
                    "FORMCYCLE-Export enthält keinen Titel - bitte zusätzlich einen Titel angeben."
                )
            felder = parse_formcycle_export(rohdaten)
            return FormularInput(titel=endgueltiger_titel, leika_nummer=leika_nummer or None, felder=felder)
        return FormularInput.model_validate(rohdaten)

    if dateiname_lower.endswith((".csv", ".xlsx")):
        if not titel:
            raise FormularImportError(
                "Für CSV/XLSX-Uploads muss zusätzlich ein Formular-Titel angegeben werden."
            )
        felder = parse_feld_datei(inhalt, dateiname)
        return FormularInput(titel=titel, leika_nummer=leika_nummer or None, felder=felder)

    raise FormularImportError(
        f"Nicht unterstütztes Dateiformat: '{dateiname}'. Erlaubt sind .json, .csv und .xlsx."
    )
